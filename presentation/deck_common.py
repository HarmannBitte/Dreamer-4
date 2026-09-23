"""Shared visual system and helpers for the Dreamer 4 decks (python-pptx). Used by build_deck.py (deep dive) and build_final.py (final talk).
All numbers come from Hafner, Yan, Lillicrap, "Training Agents Inside of Scalable World Models" (arXiv 2509.24527)
as digested in ../dreamer4_research_report.md; figures are from the paper/project page (credited on each slide)."""
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn
from lxml import etree
from PIL import Image, ImageEnhance
import os

A = "assets/"
MEDIA = "/home/user/dreamer4_resources/media/"   # archived official clips (teaser embedded on slide 4)
# ---- visual system: one dark ink, one accent, greys and hairlines; no fills, shadows or rounded shapes
NAVY = RGBColor(0x14, 0x21, 0x3D); BLUE = RGBColor(0x2A, 0x56, 0xD6); DARK = RGBColor(0x1F, 0x29, 0x37)
GREY = RGBColor(0x6B, 0x72, 0x80); MUTED = RGBColor(0x9A, 0xA3, 0xB2); RULE = RGBColor(0xD3, 0xD8, 0xE2)
PANEL = RGBColor(0xF3, 0xF5, 0xF8); WHITE = RGBColor(0xFF, 0xFF, 0xFF); SKY = RGBColor(0x9D, 0xB8, 0xF5)
GREEN = RGBColor(0x2E, 0x8B, 0x57); AMBER = RGBColor(0xC9, 0x7A, 0x0A)
LIGHT = PANEL; PALE_BLUE = PANEL; TEAL = SKY; PURPLE = BLUE; RED = AMBER   # legacy names
FONT = "Calibri"
W, H = Inches(13.333), Inches(7.5)

prs = Presentation(); prs.slide_width = W; prs.slide_height = H
BLANK = prs.slide_layouts[6]
FOOTER = "Dreamer 4  ·  Hafner, Yan, Lillicrap (Google DeepMind)  ·  arXiv 2509.24527"
state = {"n": 0, "part": None, "titles": []}

# ---------------------------------------------------------------- helpers
def _font(run_or_p, size, bold=False, color=DARK, italic=False, name=FONT):
    f = run_or_p.font; f.size = Pt(size); f.bold = bold; f.italic = italic; f.name = name; f.color.rgb = color

import re
_MARK = re.compile(r"(?<!\w)\*\*(?=\S)(.+?)(?<=\S)\*\*(?!\w)|(?<!\w)\*(?=\S)(.+?)(?<=\S)\*(?!\w)")

def add_runs(p, text, size, color=DARK, bold=False, italic=False, bold_color=None):
    """Add runs to paragraph p, honouring **bold** and *italic* inline markup (author asterisks like 'Hafner*' are left alone)."""
    pos = 0
    for m in _MARK.finditer(text):
        if m.start() > pos:
            r = p.add_run(); r.text = text[pos:m.start()]; _font(r, size, bold, color, italic)
        if m.group(1) is not None:
            r = p.add_run(); r.text = m.group(1); _font(r, size, True, bold_color or color, italic)
        else:
            r = p.add_run(); r.text = m.group(2); _font(r, size, bold, color, True)
        pos = m.end()
    if pos < len(text) or not text:
        r = p.add_run(); r.text = text[pos:]; _font(r, size, bold, color, italic)

def add_text(slide, left, top, width, height, text, size=18, bold=False, color=DARK, align=PP_ALIGN.LEFT,
             anchor=MSO_ANCHOR.TOP, italic=False, line_spacing=1.05, space=0, spc=None, bold_color=None):
    tb = slide.shapes.add_textbox(left, top, width, height); tf = tb.text_frame; tf.word_wrap = True
    tf.vertical_anchor = anchor; tf.margin_left = tf.margin_right = Inches(0.05); tf.margin_top = tf.margin_bottom = Inches(0.02)
    lines = text if isinstance(text, list) else [text]
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align; p.line_spacing = line_spacing
        if space: p.space_after = Pt(space)
        add_runs(p, line, size, color, bold, italic, bold_color)
        if spc is not None:
            for r in p.runs: r._r.get_or_add_rPr().set("spc", str(int(spc * 100)))
    return tb

def _bullet(p, level, char="•", color=MUTED):
    pPr = p._p.get_or_add_pPr()
    mar = Inches(0.26 + 0.3 * level); pPr.set("marL", str(int(mar))); pPr.set("indent", str(int(-Inches(0.22))))
    for tag in ("a:buClr", "a:buNone", "a:buChar", "a:buAutoNum", "a:buFont"):
        for el in pPr.findall(qn(tag)): pPr.remove(el)
    bc = etree.SubElement(pPr, qn("a:buClr")); c = etree.SubElement(bc, qn("a:srgbClr")); c.set("val", str(color))
    bf = etree.SubElement(pPr, qn("a:buFont")); bf.set("typeface", "Arial")
    ch = etree.SubElement(pPr, qn("a:buChar")); ch.set("char", char)

def add_bullets(slide, left, top, width, height, items, size=18, color=DARK, space=6, bold_lead=True):
    """items: list of str or (str, level). '**text**' segments are bold (lead-ins render in ink colour)."""
    tb = slide.shapes.add_textbox(left, top, width, height); tf = tb.text_frame; tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(0.05)
    for i, it in enumerate(items):
        txt, lvl = (it, 0) if isinstance(it, str) else it
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(space); p.line_spacing = 1.08
        sz = size if lvl == 0 else size - 2
        add_runs(p, txt, sz, color, bold_color=NAVY)
        _bullet(p, lvl, "•" if lvl == 0 else "–")
    return tb

def fit_picture(slide, path, left, top, max_w, max_h, align="center"):
    im = Image.open(path); ar = im.width / im.height
    w = max_w; h = int(w / ar)
    if h > max_h: h = max_h; w = int(h * ar)
    dx = (max_w - w) // 2 if align == "center" else 0
    return slide.shapes.add_picture(path, left + dx, top, width=w, height=h)

def caption(slide, left, top, width, text, size=10.5):
    return add_text(slide, left, top, width, Inches(0.5), text, size=size, color=GREY, italic=True)

def hline(slide, left, top, width, color=RULE, pt=0.75):
    r = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, Pt(pt)); r.fill.solid(); r.fill.fore_color.rgb = color; r.line.fill.background(); return r

def vline(slide, left, top, height, color=RULE, pt=0.75):
    r = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, Pt(pt), height); r.fill.solid(); r.fill.fore_color.rgb = color; r.line.fill.background(); return r

def kicker(slide, left, top, width, text, color=GREY, size=10, align=PP_ALIGN.LEFT):
    """Small upper-case, letter-spaced label."""
    return add_text(slide, left, top, width, Inches(0.3), text.upper(), size=size, bold=True, color=color, spc=1.2, align=align)

def chrome(slide, title, subtitle=None, dark=False):
    state["n"] += 1; state["titles"].append(title)
    if not dark:
        add_text(slide, Inches(0.5), Inches(0.38), Inches(12.3), Inches(0.75), title, size=26, bold=True, color=NAVY)
        if subtitle: add_text(slide, Inches(0.52), Inches(1.0), Inches(12.3), Inches(0.45), subtitle, size=14, color=GREY)
        hline(slide, Inches(0.5), Inches(6.98), Inches(12.33))
        add_text(slide, Inches(0.5), Inches(7.04), Inches(8), Inches(0.3), FOOTER, size=9, color=MUTED)
        if state["part"]: kicker(slide, Inches(8.0), Inches(7.05), Inches(4.0), state["part"], color=MUTED, size=8.5, align=PP_ALIGN.RIGHT)
    add_text(slide, Inches(12.2), Inches(7.03), Inches(0.63), Inches(0.3), str(state["n"]), size=10, bold=True, color=(SKY if dark else GREY), align=PP_ALIGN.RIGHT)

def notes(slide, text):
    slide.notes_slide.notes_text_frame.text = text

def new_slide(title, subtitle=None):
    ov = state.pop("override", None)                      # set by titled(): lets a deck reuse a slide under a different title/subtitle
    if ov: title = ov[0] or title; subtitle = ov[1] if ov[1] is not None else subtitle
    s = prs.slides.add_slide(BLANK); chrome(s, title, subtitle); return s

def titled(fn, title=None, subtitle=None, **kw):
    """Call slide function fn with its title and/or subtitle replaced (subtitle="" removes it)."""
    state["override"] = (title, subtitle); return fn(**kw)

def section(title, sub, n):
    s = prs.slides.add_slide(BLANK)
    bg = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, W, H); bg.fill.solid(); bg.fill.fore_color.rgb = NAVY; bg.line.fill.background()
    add_text(s, Inches(7.6), Inches(0.9), Inches(5.3), Inches(4.2), f"{n:02d}", size=260, bold=True, color=RGBColor(0x1D, 0x2D, 0x52), align=PP_ALIGN.RIGHT)
    kicker(s, Inches(0.9), Inches(2.55), Inches(6), f"Part {n}", color=SKY, size=13)
    hline(s, Inches(0.9), Inches(2.98), Inches(0.9), color=SKY, pt=1.5)
    add_text(s, Inches(0.9), Inches(3.15), Inches(9), Inches(1.3), title, size=44, bold=True, color=WHITE)
    add_text(s, Inches(0.9), Inches(4.35), Inches(8.5), Inches(1.4), sub, size=18, color=RGBColor(0xC9, 0xD3, 0xE6), line_spacing=1.15)
    state["part"] = f"Part {n} · {title}"
    chrome(s, title, dark=True); state["titles"][-1] = f"Part {n} — {title}"; return s

def _borders(cell, top=None, bottom=None):
    """Booktabs-style borders: (color, pt) for top/bottom, no vertical rules. Elements are inserted in schema order."""
    tcPr = cell._tc.get_or_add_tcPr()
    for tag in ("a:lnL", "a:lnR", "a:lnT", "a:lnB"):
        for el in tcPr.findall(qn(tag)): tcPr.remove(el)
    for i, (tag, spec) in enumerate((("a:lnL", None), ("a:lnR", None), ("a:lnT", top), ("a:lnB", bottom))):
        ln = etree.Element(qn(tag))
        if spec:
            ln.set("w", str(int(Pt(spec[1])))); ln.set("cap", "flat"); ln.set("cmpd", "sng"); ln.set("algn", "ctr")
            sf = etree.SubElement(ln, qn("a:solidFill")); c = etree.SubElement(sf, qn("a:srgbClr")); c.set("val", str(spec[0]))
            d = etree.SubElement(ln, qn("a:prstDash")); d.set("val", "solid")
        else:
            ln.set("w", "0"); etree.SubElement(ln, qn("a:noFill"))
        tcPr.insert(i, ln)

def add_table(slide, left, top, width, rows, col_widths=None, font=12, header_fill=None, highlight_col=None, highlight_row=None, row_h=0.36):
    nr, nc = len(rows), len(rows[0])
    shp = slide.shapes.add_table(nr, nc, left, top, width, Inches(row_h * nr)); t = shp.table
    tblPr = t._tbl.tblPr; tblPr.set("firstRow", "0"); tblPr.set("bandRow", "0")
    for el in tblPr.findall(qn("a:tableStyleId")): tblPr.remove(el)
    if col_widths:
        for i, cw in enumerate(col_widths): t.columns[i].width = Inches(cw)
    if nr > 1 and row_h > 0.45:
        t.rows[0].height = Inches(0.42)
    for r in range(nr):
        for c in range(nc):
            cell = t.cell(r, c); cell.text = ""; cell.margin_left = cell.margin_right = Inches(0.07); cell.margin_top = cell.margin_bottom = Inches(0.04)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            p = cell.text_frame.paragraphs[0]; txt = str(rows[r][c])
            hdr = (r == 0); hl = (highlight_row == r) or (highlight_col == c and r > 0)
            add_runs(p, txt, font if not hdr else font - 0.5, color=NAVY if hdr else (BLUE if (highlight_col == c and r > 0) else DARK), bold=hdr or hl)
            p.alignment = PP_ALIGN.LEFT if (c == 0 or len(txt) > 28) else PP_ALIGN.CENTER
            if highlight_row == r: cell.fill.solid(); cell.fill.fore_color.rgb = PANEL
            else: cell.fill.background()
            if hdr: _borders(cell, top=(NAVY, 1.0), bottom=(NAVY, 1.0))
            elif r == nr - 1: _borders(cell, bottom=(NAVY, 1.0))
            else: _borders(cell, bottom=(RULE, 0.5))
    return t

def note(slide, left, top, width, height, title, body, size=12.5, accent=BLUE, **_):
    """Margin note: thin accent rule on the left, small-caps label, body text. Replaces the old filled/rounded box."""
    body = body if isinstance(body, list) else [body]
    cpl = max(20, int((width - Inches(0.22)) / Inches(1) * 72 / (size * 0.46)))
    est = Inches(0.42 + sum(-(-len(re.sub(r"\*", "", b)) // cpl) for b in body) * size * 1.1 * 1.2 / 72 + len(body) * 5 / 72)
    vline(slide, left, top + Inches(0.03), min(height, est) - Inches(0.06), color=accent, pt=2)
    kicker(slide, left + Inches(0.2), top, width - Inches(0.22), title, color=NAVY, size=10)
    if body and all(b.startswith("• ") for b in body):
        add_bullets(slide, left + Inches(0.14), top + Inches(0.36), width - Inches(0.16), height - Inches(0.4), [b[2:] for b in body], size=size, space=4)
    else:
        add_text(slide, left + Inches(0.2), top + Inches(0.36), width - Inches(0.22), height - Inches(0.4), body, size=size, color=DARK, line_spacing=1.1, space=5)
box = note

def columns(slide, top, height, items, left=Inches(0.5), width=Inches(12.33), gap=Inches(0.45), start=1, size=13, head_size=15,
            numerals=True, num_label=lambda i: f"{i:02d}", arrows=False, bullets=False, rule_colors=None, head_h=0.42):
    """Flat column layout: optional accent numeral, thin rule, heading, body. items = [(heading, body-str-or-list), ...]."""
    n = len(items); cw = int((width - gap * (n - 1)) / n)
    for i, (head, body) in enumerate(items):
        x = left + i * (cw + gap); y = top
        if numerals:
            add_text(slide, x, y, cw, Inches(0.45), num_label(start + i), size=20, bold=True, color=BLUE)
            if arrows and i < n - 1: add_text(slide, x + cw, y + Inches(0.02), gap, Inches(0.45), "→", size=18, color=MUTED, align=PP_ALIGN.CENTER)
            y += Inches(0.48)
        hline(slide, x, y, cw, color=(rule_colors[i] if rule_colors else NAVY), pt=1.0)
        add_text(slide, x, y + Inches(0.08), cw, Inches(head_h), head, size=head_size, bold=True, color=NAVY)
        by = y + Inches(0.12) + Inches(head_h); bh = top + height - by
        body = body if isinstance(body, list) else [body]
        if bullets: add_bullets(slide, x - Inches(0.05), by, cw + Inches(0.05), bh, body, size=size, space=5)
        else: add_text(slide, x, by, cw, bh, body, size=size, color=DARK, line_spacing=1.1, space=5)

def numbered(slide, left, top, width, items, size=16, gap=0.22, num_w=0.55, head_w=None):
    """Numbered list with accent numerals in a narrow column. items = str | (head, body). Returns bottom y."""
    y = top
    for i, it in enumerate(items):
        head, body = (None, it) if isinstance(it, str) else it
        add_text(slide, left, y - Inches(0.03), Inches(num_w), Inches(0.5), f"{i + 1:02d}", size=size + 2, bold=True, color=BLUE)
        tx = left + Inches(num_w); tw = width - Inches(num_w)
        if head and head_w:
            add_text(slide, tx, y, Inches(head_w), Inches(0.5), head, size=size, bold=True, color=NAVY); tx += Inches(head_w); tw -= Inches(head_w)
            txt = body
        else:
            txt = (f"**{head}** {body}" if head else body)
        cpl = max(20, int(tw / Inches(1) * 72 / (size * 0.46)))      # rough chars per line
        lines = max(1, -(-len(re.sub(r"\*", "", txt)) // cpl))
        add_text(slide, tx, y, tw, Inches(0.3 + lines * size * 1.25 / 72), txt, size=size, color=DARK, line_spacing=1.1, bold_color=NAVY)
        y += Inches(lines * size * 1.25 / 72 + gap)
    return y

def arrow(slide, left, top, width=Inches(0.45), height=Inches(0.5), color=GREY):
    return add_text(slide, left, top, width, height, "→", size=18, color=MUTED, align=PP_ALIGN.CENTER)


def caption_links(slide, left, top, width, parts, size=10.5, color=GREY, link_color=BLUE):
    """Italic caption made of (text, url-or-None) parts; URL parts become blue hyperlinks."""
    tb = slide.shapes.add_textbox(left, top, width, Inches(0.5)); tf = tb.text_frame; tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(0.05); tf.margin_top = tf.margin_bottom = Inches(0.02)
    p = tf.paragraphs[0]
    for text, url in parts:
        r = p.add_run(); r.text = text; _font(r, size, bool(url), link_color if url else color, italic=not url)
        if url: r.hyperlink.address = url
    return tb

def click_through(shape, url):
    shape.click_action.hyperlink.address = url; return shape

def jpeg(path, max_w=1600, q=85):
    """Downscale/convert an asset to JPEG to keep the deck small; returns new path."""
    out = path.rsplit(".", 1)[0] + "_s.jpg"
    if os.path.exists(out): return out
    im = Image.open(path).convert("RGB")
    if im.width > max_w: im = im.resize((max_w, int(im.height * max_w / im.width)), Image.LANCZOS)
    im.save(out, quality=q, optimize=True); return out

YT = "https://www.youtube.com/watch?v=oDlBtTcX0g0"; PROJECT = "https://danijar.com/project/dreamer4"

def finish(default_out, notes_md, heading):
    """Save the deck (path overridable with OUT=...) and export the speaker notes as Markdown."""
    out = os.environ.get("OUT", default_out); prs.save(out)
    print("saved", out, os.path.getsize(out) // 1024, "KB,", len(prs.slides), "slides")
    with open(notes_md, "w") as f:
        f.write(f"# {heading}\n\n")
        for i, (sl, t) in enumerate(zip(prs.slides, state["titles"]), 1):
            txt = sl.notes_slide.notes_text_frame.text.strip() if sl.has_notes_slide else ""
            f.write(f"## {i}. {t}\n\n{txt or '_(no notes)_'}\n\n")
    print("notes written for", len(state["titles"]), "slides")

